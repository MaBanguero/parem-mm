"""Rellena la plantilla oficial de la matriz PARE.MM preservando su formato.

La plantilla tiene la cabecera en la fila 5 (B..N) y UNA fila de datos por
municipio (filas 6/7/8), con el municipio en la columna D. Las métricas van en
las columnas I..M (orden: pruebas embarazo, gestantes captadas, control
prenatal, puerperio, RN). El resto (ESE, resolución, EBS, territorios,
microterritorios) queda intacto.
"""

import io

import openpyxl

COL_MUNICIPIO = 4          # columna D
COLS_METRICAS = ["I", "J", "K", "L", "M"]
FILAS_DATOS = (6, 7, 8)    # Puerto Tejada, Padilla, Villa Rica


def _norm_mun(v):
    import unicodedata

    if not v:
        return None
    s = unicodedata.normalize("NFD", str(v)).strip().upper()
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s


def build(template_path: str, metrics_by_municipio: dict) -> io.BytesIO:
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    # Construir índice municipio -> métricas (normalizado)
    idx = {_norm_mun(m): v for m, v in metrics_by_municipio.items()}

    for row in FILAS_DATOS:
        mun = _norm_mun(ws.cell(row=row, column=COL_MUNICIPIO).value)
        datos = idx.get(mun)
        if datos is None:
            continue
        for i, col in enumerate(COLS_METRICAS):
            ws[f"{col}{row}"] = int(datos[col])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
